"""Build the candidate roster template."""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

ARIAL = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")     # a CV would fill this
ONLY_HERE_HEAD = PatternFill("solid", fgColor="375623")  # nothing else holds it
CV_FILL = PatternFill("solid", fgColor="DDEBF7")      # a CV would fill this
ONLY_HERE_FILL = PatternFill("solid", fgColor="E2EFDA")  # nothing else holds it
EXAMPLE_FILL = PatternFill("solid", fgColor="FFF2CC")
REQUIRED_FILL = PatternFill("solid", fgColor="FCE4D6")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# name, width, from_cv, example
COLUMNS = [
    ("full_name", 24, True, "Thandiwe Mokoena"),
    ("email", 28, True, "t.mokoena@example.co.za"),
    ("phone", 18, True, "082 555 0143"),
    ("location", 16, True, "Johannesburg"),
    ("current_role", 26, True, "ERP Consultant"),
    ("additional_roles", 30, True, "Business Analyst | Systems Analyst"),
    ("years_experience", 18, True, 9),
    ("designated_group", 20, True, "African Female"),
    ("qualifications", 38, True, "BCom Information Systems | Matric"),
    ("certifications", 38, True, "Microsoft Certified: Dynamics 365 F&O Apps Developer Associate"),
    ("technical_skills", 38, True, "Microsoft Dynamics 365 | SQL | X++"),
    ("skills", 34, True, "Business Process Analysis | UAT"),
    ("sectors", 26, True, "Public Sector | Finance"),
    ("languages", 22, True, "English | Setswana"),
    ("availability", 16, False, "available"),
    ("available_from", 18, False, ""),
    ("status", 14, False, "active"),
    ("resource_categories", 30, False, "ERP | Business Analysis"),
    ("notes", 40, False, "Strong on D365 Finance. Prefers Gauteng-based work."),
]

DATE_COLUMNS = {"available_from"}

wb = Workbook()

# ------------------------------------------------------------- Candidates ---
ws = wb.active
ws.title = "Candidates"

for idx, (name, width, from_cv, example) in enumerate(COLUMNS, start=1):
    letter = get_column_letter(idx)
    ws.column_dimensions[letter].width = width

    # The header colour says where each column's answer really comes from,
    # because that is the whole point: fourteen of these a CV answers better.
    # Carried by the colour rather than a label row, since a row saying
    # "on the CV" nineteen times becomes a data row the moment this is saved
    # as CSV, and the importer would try to read it as a person.
    head = ws.cell(row=1, column=idx, value=name)
    head.font = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
    head.fill = HEADER_FILL if from_cv else ONLY_HERE_HEAD
    head.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    head.border = BORDER

    cell = ws.cell(row=2, column=idx, value=example)
    cell.font = Font(name=ARIAL, size=10, italic=True)
    cell.fill = EXAMPLE_FILL
    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    cell.border = BORDER
    if name in DATE_COLUMNS:
        cell.number_format = "@"

ws.row_dimensions[1].height = 30
ws.row_dimensions[2].height = 34
ws.freeze_panes = "A3"

def dropdown(col_name, values, message):
    col = get_column_letter([c[0] for c in COLUMNS].index(col_name) + 1)
    rule = DataValidation(type="list", formula1=f'"{values}"', allow_blank=True,
                          showErrorMessage=True)
    rule.error = message
    rule.errorTitle = "Not a recognised value"
    ws.add_data_validation(rule)
    rule.add(f"{col}2:{col}500")

dropdown("availability", "available,notice_period,unavailable",
         "Use one of: available, notice_period, unavailable")
dropdown("status", "active,inactive", "Use one of: active, inactive. "
         "Placed is set by the system when somebody is put on a contract.")

for name in DATE_COLUMNS:
    col = get_column_letter([c[0] for c in COLUMNS].index(name) + 1)
    for row in range(3, 501):
        ws[f"{col}{row}"].number_format = "@"

# Deliberately NOT a cell below the table. Excel exports every cell in the used
# range, so a note written under the data becomes a row whose name is the note.
# A comment carries the same words and never reaches the CSV.
ws["A1"].comment = Comment(
    "Row 2 is an example. Delete it before sending this back, and put one person "
    "per row from row 2 down.\n\n"
    "Header colours: dark blue means a CV already answers this column better. "
    "Green means nothing but this sheet holds it.\n\n"
    "Do not type anything below or beside the table: saving as CSV picks up every "
    "cell on this sheet, and a stray note becomes a person.\n\n"
    "Read the 'How to fill this in' tab first, especially the part about people "
    "who already have a TiPP Focus CV.",
    "Resource Planning",
)
ws["A1"].comment.width = 340
ws["A1"].comment.height = 210

# ------------------------------------------------------ How to fill it in ---
guide = wb.create_sheet("How to fill this in")
guide.column_dimensions["A"].width = 26
guide.column_dimensions["B"].width = 14
guide.column_dimensions["C"].width = 60
guide.column_dimensions["D"].width = 38

title = guide.cell(row=1, column=1, value="Candidate roster: what goes in each column")
title.font = Font(name=ARIAL, size=14, bold=True, color="1F3864")
guide.merge_cells("A1:D1")

r = 2
warn_head = guide.cell(row=r, column=1, value="Read this before filling anything in")
warn_head.font = Font(name=ARIAL, size=12, bold=True, color="C00000")
guide.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
r += 1

warn = guide.cell(row=r, column=1, value=(
    "If somebody already has a CV on the TiPP Focus template, upload the CV instead of typing "
    "them into this sheet. The system reads that template accurately and gets their full work "
    "history and education, which this sheet has no room for and which matters on a bid. "
    "Retyping a CV here produces a thinner record, not a faster one.\n\n"
    "This sheet earns its keep in two cases. People who have no CV yet, so there is nothing to "
    "upload. And the last five columns, which no CV carries: whether somebody is free, when they "
    "come free, whether they are still on the roster, which practice areas they sit in, and "
    "anything a resourcing conversation needs to know. Those can be filled in here for people "
    "whose CVs are already loaded, and left blank for everyone else."
))
warn.font = Font(name=ARIAL, size=10)
warn.alignment = Alignment(wrap_text=True, vertical="top")
guide.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
guide.row_dimensions[r].height = 118
r += 1

legend = guide.cell(row=r, column=1, value="Header colours on the Candidates tab")
legend.font = Font(name=ARIAL, size=10, bold=True)
lg1 = guide.cell(row=r, column=2, value="on the CV")
lg1.font = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
lg1.fill = HEADER_FILL
lg1.alignment = Alignment(horizontal="center")
lg2 = guide.cell(row=r, column=3, value="only here")
lg2.font = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
lg2.fill = ONLY_HERE_HEAD
lg2.alignment = Alignment(horizontal="left", indent=1)
r += 2

heads = ["Column", "Comes from", "What it is", "Example"]
for idx, text in enumerate(heads, start=1):
    c = guide.cell(row=r, column=idx, value=text)
    c.font = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
    c.fill = HEADER_FILL
    c.border = BORDER
    c.alignment = Alignment(horizontal="left", vertical="center")
r += 1

GUIDE = [
    ("full_name", "CV", "Required. The only column that cannot be blank. As it should read on a submitted CV.",
     "Thandiwe Mokoena"),
    ("email", "CV", "Work or personal address.", "t.mokoena@example.co.za"),
    ("phone", "CV", "Any format.", "082 555 0143"),
    ("location", "CV", "Where they are based.", "Johannesburg"),
    ("current_role", "CV",
     "Their main role, spelled the way a tender would name it. This is the one that decides "
     "matching, so 'Business Analyst', never 'BA'.",
     "ERP Consultant"),
    ("additional_roles", "CV",
     "Other roles they can be put forward for, separated by a pipe. Matching takes whichever "
     "scores best, so this genuinely widens what they can be bid on.",
     "Business Analyst | Systems Analyst"),
    ("years_experience", "CV", "Whole years. Used against a seat's minimum.", "9"),
    ("designated_group", "CV",
     "Employment equity group, as it should read on a submitted CV. Matters on a BEE-scored bid.",
     "African Female"),
    ("qualifications", "CV", "Degrees and diplomas, separated by a pipe.",
     "BCom Information Systems | Matric"),
    ("certifications", "CV", "Certifications, separated by a pipe. Worth 20 of the 100 matching points.",
     "Microsoft Certified: Dynamics 365 F&O Apps Developer Associate"),
    ("technical_skills", "CV", "Tools, platforms and languages, separated by a pipe.",
     "Microsoft Dynamics 365 | SQL | X++"),
    ("skills", "CV", "Professional and domain skills, separated by a pipe.",
     "Business Process Analysis | UAT"),
    ("sectors", "CV", "Sectors they have worked in, separated by a pipe.", "Public Sector | Finance"),
    ("languages", "CV", "Separated by a pipe.", "English | Setswana"),
    ("availability", "Only here",
     "One of: available, notice_period, unavailable. Pick from the dropdown. Blank is read as available.",
     "available"),
    ("available_from", "Only here",
     "The date they next come free, if they are committed now. Leave blank if free today. "
     "Year, then month, then day.",
     "2027-03-01"),
    ("status", "Only here",
     "One of: active or inactive. Blank is read as active. Inactive takes somebody out of "
     "matching without deleting them. Placed is not typed here: the system sets it when "
     "somebody is actually put on a contract.",
     "active"),
    ("resource_categories", "Only here",
     "Practice areas, separated by a pipe. Used for the bench view and reporting. "
     "One of: ERP, Enterprise Architecture (EA), Software Development, Data & BI, Cloud & DevOps, "
     "Cybersecurity, Project & Programme Management, Business Analysis, Finance & Accounting, "
     "Construction & Engineering.",
     "ERP | Business Analysis"),
    ("notes", "Only here", "Anything a resourcing conversation needs. Free text.",
     "Strong on D365 Finance. Prefers Gauteng-based work."),
]

for name, source, meaning, example in GUIDE:
    a = guide.cell(row=r, column=1, value=name)
    a.font = Font(name=ARIAL, size=10, bold=True)
    b = guide.cell(row=r, column=2, value=source)
    b.font = Font(name=ARIAL, size=10)
    b.fill = CV_FILL if source == "CV" else ONLY_HERE_FILL
    if name == "full_name":
        b.fill = REQUIRED_FILL
    c = guide.cell(row=r, column=3, value=meaning)
    c.font = Font(name=ARIAL, size=10)
    d = guide.cell(row=r, column=4, value=example)
    d.font = Font(name=ARIAL, size=10, italic=True)
    for col in range(1, 5):
        cell = guide.cell(row=r, column=col)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.border = BORDER
    guide.row_dimensions[r].height = 60 if len(meaning) > 110 else (44 if len(meaning) > 60 else 30)
    r += 1

r += 1
rules_head = guide.cell(row=r, column=1, value="Three things that cause trouble")
rules_head.font = Font(name=ARIAL, size=12, bold=True, color="C00000")
guide.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
r += 1

RULES = [
    ("Role names decide everything",
     "current_role is compared against the role on a tender seat. It scores full marks for an "
     "exact match and nothing at all for a spelling the system does not know, and role is 35 of "
     "the 100 matching points. Somebody entered as 'Snr BA' will never come up as a strong match "
     "for a Business Analyst seat, and nothing on screen will explain why. Write the role out."),
    ("Dates must read 2027-03-01",
     "Year, then month, then day. Only available_from is a date, and only for people who are "
     "committed right now."),
    ("Use a pipe between list items",
     "Every list column separates with a pipe, not a comma, so a qualification like "
     "'BCom Information Systems, Cum Laude' stays as one thing."),
]
for heading, detail in RULES:
    h = guide.cell(row=r, column=1, value=heading)
    h.font = Font(name=ARIAL, size=10, bold=True)
    h.alignment = Alignment(wrap_text=True, vertical="top")
    d = guide.cell(row=r, column=3, value=detail)
    d.font = Font(name=ARIAL, size=10)
    d.alignment = Alignment(wrap_text=True, vertical="top")
    guide.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
    guide.row_dimensions[r].height = 68
    r += 1

r += 1
send = guide.cell(row=r, column=1, value=(
    "When it is filled in: save as CSV (File, Save As, CSV UTF-8) and send that. "
    "It is read with  npx tsx scripts/import-candidates.ts <file.csv>  which prints back "
    "exactly what it would do and writes nothing. Only a second run with --apply writes. "
    "For anybody who already has a CV on file, it fills in a blank column and never "
    "overwrites what the CV said, and the report names every column it left alone."
))
send.font = Font(name=ARIAL, size=10, bold=True)
send.alignment = Alignment(wrap_text=True, vertical="top")
guide.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
guide.row_dimensions[r].height = 30

wb.save(sys.argv[1])
print("wrote", sys.argv[1])
